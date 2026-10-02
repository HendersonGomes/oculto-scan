"""Clean copy: the original stays byte-for-byte, formulas stay unless asked."""

import hashlib
import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook

from oculto_scan.cli import main
from oculto_scan.gui_logic import (
    CLEAN_BUTTON,
    CLEAN_FORCE,
    CLEAN_HIDDEN,
    CLEAN_MACROS,
    Session,
    clean_button_enabled,
    finish_clean_session,
)
from oculto_scan.zipsafe import open_office_bytes
from tests.workbook_factory import build_workbook, make_cpf

ROOT = Path(__file__).resolve().parents[1]
CPF = make_cpf("529982247")
_UNC = "\\\\SERVIDOR-OBRAS\\propostas"


def _rich(path: Path) -> None:
    build_workbook(
        path,
        [
            {
                "name": "Proposta",
                "cells": [
                    {"ref": "A1", "value": "Item"},
                    {"ref": "A2", "value": 10},
                    {"ref": "B1", "value": "CPF"},
                    {"ref": "A3", "value": "Funcionário"},
                    {"ref": "B3", "value": CPF},
                    {"ref": "C2", "formula": "A2*1.1", "value": 11},
                    {"ref": "D2", "formula": "Custos!B2*1.35", "value": 135},
                    {"ref": "E2", "formula": "'[Custos.xlsx]Custos'!A1", "value": 50},
                    {"ref": "G2", "formula": "'[Custos.xlsx]Custos'!B9"},
                    {"ref": "F3", "value": "senha: sintetica-123"},
                    {"ref": "A8", "value": "marcador-oculto"},
                    {"ref": "A9", "value": "marcador-visivel"},
                ],
                "hidden_rows": [8],
                "comments": [{"ref": "A1", "author": "Ana Silva", "text": "conferir margem"}],
            },
            {"name": "Custos", "state": "hidden", "cells": [{"ref": "B2", "value": 100}]},
            {
                "name": "Medicao",
                "cells": [
                    {"ref": "B1", "value": 10},
                    {"ref": "B2", "formula": "B1*1.1", "value": 11},
                ],
            },
        ],
        metadata={
            "creator": "Ana Silva",
            "lastModifiedBy": "Ana Silva",
            "company": "Construtora Exemplo",
            "created": "2026-01-01T00:00:00Z",
            "modified": "2026-02-01T00:00:00Z",
        },
        external_target="file:///C:/Obras/custos.xlsx",
        hyperlink={"sheet": "Proposta", "target": "file:///C:/Obras/anexo.pdf"},
        defined_names=[
            {"name": "CustoExterno", "formula": "'[Custos.xlsx]Custos'!A1"},
            {"name": "Margem", "formula": "Proposta!A2"},
            {"name": "Oculto", "formula": "Custos!B2"},
        ],
        custom={"Servidor": _UNC, "Obra": "Ponte"},
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _names(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return archive.namelist()


def _xml(path: Path) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                chunks.append(archive.read(name).decode("utf-8"))
    return "\n".join(chunks)


def test_original_stays_and_local_formula_is_kept(tmp_path, capsys):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    digest = _sha256(path)
    with zipfile.ZipFile(path) as archive:
        untouched = archive.read("xl/worksheets/sheet3.xml")
    code = main([str(path), "--limpar", "--no-color"])
    text = capsys.readouterr().out
    dest = tmp_path / "proposta-limpa.xlsx"
    assert code == 1
    assert dest.is_file()
    assert _sha256(path) == digest
    assert path.read_bytes() != dest.read_bytes()
    text_xml = _xml(dest)
    assert "A2*1.1" in text_xml
    assert "Custos!B2*1.35" in text_xml
    assert "B1*1.1" in text_xml
    assert "[Custos.xlsx]" not in text_xml
    assert "C:/Obras/custos.xlsx" not in text_xml
    assert "anexo.pdf" in text_xml
    assert "Ana Silva" not in text_xml
    assert "SERVIDOR-OBRAS" not in text_xml
    assert "Construtora Exemplo" not in text_xml
    assert "Ponte" in text_xml
    assert CPF in text_xml
    assert "sintetica-123" in text_xml
    assert "marcador-oculto" in text_xml
    assert "marcador-visivel" in text_xml
    assert not any("comments" in name.casefold() for name in _names(dest))
    with zipfile.ZipFile(dest) as archive:
        assert archive.read("xl/worksheets/sheet3.xml") == untouched
        workbook = archive.read("xl/workbook.xml")
    assert b"CustoExterno" not in workbook
    assert b"Margem" in workbook
    assert b"Oculto" in workbook
    assert b'state="hidden"' in workbook
    archive = open_office_bytes(dest.read_bytes())
    assert archive.testzip() is None
    archive.close()
    book = load_workbook(dest)
    assert "Custos" in book.sheetnames
    assert "A2*1.1" in str(book["Proposta"]["C2"].value)
    assert "Custos!B2*1.35" in str(book["Proposta"]["D2"].value)
    assert int(book["Proposta"]["E2"].value) == 50
    assert book["Proposta"]["B3"].value == CPF
    assert "O original não foi alterado" in text
    assert "Removido:" in text
    assert "Ainda na cópia:" in text
    assert "revisão humana" in text
    assert "--remover-ocultas" in text
    assert "sem valor em cache" in text
    assert "Antes:" in text
    assert "Depois:" in text
    assert CPF not in text
    assert "Ana Silva" not in text
    assert "SERVIDOR-OBRAS" not in text
    assert "sintetica-123" not in text


def test_hidden_removal_keeps_other_formulas_and_row_numbers(tmp_path, capsys):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    digest = _sha256(path)
    code = main([str(path), "--limpar", "--remover-ocultas", "--no-color"])
    text = capsys.readouterr().out
    dest = tmp_path / "proposta-limpa.xlsx"
    assert code == 1
    assert _sha256(path) == digest
    text_xml = _xml(dest)
    assert "A2*1.1" in text_xml
    assert "Custos!B2" not in text_xml
    assert "marcador-oculto" not in text_xml
    assert "marcador-visivel" in text_xml
    assert 'r="9"' in text_xml
    assert CPF in text_xml
    book = load_workbook(dest)
    assert "Custos" not in book.sheetnames
    assert "A2*1.1" in str(book["Proposta"]["C2"].value)
    assert int(book["Proposta"]["D2"].value) == 135
    assert book["Medicao"]["B2"].value is not None
    assert "Margem" in text_xml
    assert "Oculto" not in text_xml
    assert "esvaziadas" in text
    archive = open_office_bytes(dest.read_bytes())
    archive.close()


def test_refuses_to_overwrite_or_replace_the_original(tmp_path, capsys):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    digest = _sha256(path)
    dest = tmp_path / "proposta-limpa.xlsx"
    dest.write_bytes(b"antigo")
    code = main([str(path), "--limpar"])
    captured = capsys.readouterr()
    assert code == 2
    assert dest.read_bytes() == b"antigo"
    assert "forcar" in captured.err
    assert _sha256(path) == digest
    code = main([str(path), "--limpar", "--saida", str(path), "--forcar"])
    assert code == 2
    assert _sha256(path) == digest
    code = main([str(path), "--forcar"])
    assert code == 2
    code = main([str(path), "--limpar", "--forcar", "--no-color"])
    assert code == 1
    assert dest.read_bytes() != b"antigo"
    assert _sha256(path) == digest


def test_macros_stay_until_asked(tmp_path, capsys):
    path = tmp_path / "medicao.xlsm"
    build_workbook(
        path,
        [{"name": "Proposta", "cells": [{"ref": "A1", "value": "ok"}]}],
        vba_blob=b"nao-e-um-projeto",
    )
    digest = _sha256(path)
    code = main([str(path), "--limpar", "--no-color"])
    text = capsys.readouterr().out
    kept = tmp_path / "medicao-limpa.xlsm"
    assert code == 0
    assert kept.is_file()
    assert "xl/vbaProject.bin" in _names(kept)
    assert "--remover-macros" in text
    assert _sha256(path) == digest
    code = main([str(path), "--limpar", "--remover-macros", "--no-color"])
    capsys.readouterr()
    cleaned = tmp_path / "medicao-limpa.xlsx"
    assert code == 0
    assert cleaned.is_file()
    assert kept.is_file()
    assert "xl/vbaProject.bin" not in _names(cleaned)
    assert cleaned.suffix == ".xlsx"
    assert _sha256(path) == digest
    book = load_workbook(cleaned)
    assert book["Proposta"]["A1"].value == "ok"


def test_show_reveals_and_ci_refuses_before_writing(tmp_path, monkeypatch, capsys):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    monkeypatch.chdir(tmp_path)
    code = main([path.name, "--limpar", "--show", "--no-color"])
    text = capsys.readouterr().out
    assert code == 1
    assert CPF in text
    assert "sintetica-123" in text
    monkeypatch.setenv("CI", "true")
    other = tmp_path / "outra.xlsx"
    _rich(other)
    blocked = tmp_path / "outra-limpa.xlsx"
    code = main([str(other), "--limpar", "--show"])
    assert code == 2
    assert not blocked.exists()


def test_json_stays_masked(tmp_path, capsys):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    code = main([str(path), "--limpar", "--format", "json"])
    raw = capsys.readouterr().out
    assert code == 1
    assert '"modo": "limpar"' in raw
    assert CPF not in raw
    assert "SERVIDOR-OBRAS" not in raw
    assert "sintetica-123" not in raw


def test_window_clean_scans_the_copy(tmp_path):
    path = tmp_path / "proposta.xlsx"
    _rich(path)
    digest = _sha256(path)
    assert CLEAN_BUTTON == "Gerar cópia limpa"
    assert "ocultas" in CLEAN_HIDDEN.casefold()
    assert "macros" in CLEAN_MACROS.casefold()
    assert "já existir" in CLEAN_FORCE
    assert clean_button_enabled(Session(), mode="scan", busy=False) is False
    source = (ROOT / "src" / "oculto_scan" / "gui_window.py").read_text(encoding="utf-8")
    assert "CLEAN_BUTTON" in source
    assert "finish_clean_session" in source
    assert "result_text(" not in source
    dest, session, note = finish_clean_session(path)
    assert dest.name == "proposta-limpa.xlsx"
    assert session.ready
    assert session.names == (dest.name,)
    assert "não foi alterado" in note
    assert _sha256(path) == digest


def test_libreoffice_opens_the_copy_when_installed(tmp_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice não está instalado")
    path = tmp_path / "proposta.xlsx"
    build_workbook(
        path,
        [
            {
                "name": "Proposta",
                "cells": [
                    {"ref": "A1", "value": "Item"},
                    {"ref": "B1", "value": 12},
                    {"ref": "B2", "formula": "B1*1.1", "value": 13},
                ],
                "comments": [{"ref": "A1", "author": "Ana", "text": "nota interna"}],
            }
        ],
        metadata={"creator": "Ana Silva", "company": "Construtora Exemplo"},
    )
    code = main([str(path), "--limpar", "--no-color"])
    assert code == 0
    dest = tmp_path / "proposta-limpa.xlsx"
    out = tmp_path / "lo"
    out.mkdir()
    env = dict(os.environ)
    env["HOME"] = str(tmp_path / "home")
    (tmp_path / "home").mkdir()
    completed = subprocess.run(
        [soffice, "--headless", "--norestore", "--convert-to", "csv", "--outdir", str(out), str(dest)],
        check=False,
        timeout=120,
        capture_output=True,
        env=env,
    )
    assert completed.returncode == 0, completed.stderr.decode("utf-8", "replace")
    csv_files = list(out.glob("*.csv"))
    assert csv_files
    assert "Item" in csv_files[0].read_text(encoding="utf-8", errors="replace")

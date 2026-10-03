"""v0.2.4: where the workbook was last saved. Paths are masked and never opened."""

import base64
import os
import shutil
import subprocess
import zipfile
from pathlib import Path

import pytest
from defusedxml import ElementTree as DefusedET
from openpyxl import load_workbook as open_xlsx

from oculto_scan.analyze import analyze
from oculto_scan.cli import main
from oculto_scan.gui_logic import compare_files, scan_file, window_cards
from oculto_scan.report import render_html, render_json, render_text
from oculto_scan.savepath import mask_save_path
from oculto_scan.workbook import load_workbook
from tests.workbook_factory import build_workbook

_LOCAL = "C:/Users/fulano/Desktop"
_ONEDRIVE = "C:/Users/fulano/OneDrive - Empresa/Obra/pasta/subpasta/nomedodocumento"
_UNC = r"\\servidor\obra\pasta\arquivo.xlsx"
_ROOT = "C:/Obras/medicao/planilha.xlsx"


def _book(path: Path, **kwargs) -> None:
    build_workbook(
        path,
        [{"name": "Proposta", "cells": [{"ref": "A1", "value": "Item"}, {"ref": "B1", "value": 12}]}],
        **kwargs,
    )


def _pasta(path: Path):
    workbook = load_workbook(path)
    findings = [item for item in analyze(workbook, path.name) if item.rule == "pasta-salva"]
    return workbook, findings


def _add_parts(path: Path, extra: dict[str, str]) -> None:
    buffer = Path(str(path) + ".tmp")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(buffer, "w") as dest:
        for info in source.infolist():
            dest.writestr(info.filename, source.read(info.filename))
        for name, text in extra.items():
            dest.writestr(name, text.encode("utf-8"))
    buffer.replace(path)


def _zip_text(path: Path) -> str:
    chunks = []
    with zipfile.ZipFile(path) as archive:
        for name in archive.namelist():
            if name.endswith((".xml", ".rels")):
                chunks.append(archive.read(name).decode("utf-8"))
    return "\n".join(chunks)


def test_mask_hides_user_company_and_deep_folders():
    masked = mask_save_path(_ONEDRIVE)
    assert masked.startswith("C:/Users/")
    assert "OneDrive - Em" in masked
    assert masked.endswith("...")
    assert "fulano" not in masked
    assert "Empresa" not in masked
    local = mask_save_path(_LOCAL)
    assert local.startswith("C:/Users/")
    assert local.endswith("/Desktop")
    assert "fulano" not in local
    unc = mask_save_path(_UNC)
    assert unc.startswith("//")
    assert "servidor" not in unc
    assert "arquivo.xlsx" not in unc


def test_local_onedrive_and_unc_abs_path(tmp_path):
    local = tmp_path / "desktop.xlsx"
    cloud = tmp_path / "nuvem.xlsx"
    share = tmp_path / "rede.xlsx"
    plain = tmp_path / "simples.xlsx"
    _book(local, abs_path=_LOCAL)
    _book(cloud, abs_path=_ONEDRIVE)
    _book(share, abs_path=_UNC)
    _book(plain)
    _workbook, desktop = _pasta(local)
    assert len(desktop) == 1
    assert desktop[0].cell == "absPath"
    assert desktop[0].type_label == "Pasta onde foi salvo"
    assert desktop[0].risk == "alto"
    assert "fulano" not in (desktop[0].evidence_masked or "")
    assert "fulano" in (desktop[0].evidence_raw or "")
    assert "usuário:" in (desktop[0].evidence_masked or "")
    assert "indício" in desktop[0].message
    _workbook, drive = _pasta(cloud)
    assert drive[0].risk == "alto"
    assert "Empresa" not in (drive[0].evidence_masked or "")
    assert "Empresa" in (drive[0].evidence_raw or "")
    assert "OneDrive/empresa:" in (drive[0].evidence_masked or "")
    assert "..." in (drive[0].evidence_masked or "")
    _workbook, remote = _pasta(share)
    assert remote[0].risk == "alto"
    assert "servidor" not in (remote[0].evidence_masked or "")
    assert "servidor" in (remote[0].evidence_raw or "")
    _workbook, quiet = _pasta(plain)
    assert quiet == []


def test_backslash_path_and_medium_folder(tmp_path):
    path = tmp_path / "barras.xlsx"
    _book(path, abs_path=r"C:\Users\fulano\Desktop")
    _workbook, findings = _pasta(path)
    assert findings[0].risk == "alto"
    assert "fulano" in (findings[0].evidence_raw or "")
    medium = tmp_path / "obras.xlsx"
    _book(medium, abs_path=_ROOT)
    _workbook, findings = _pasta(medium)
    assert findings[0].risk == "medio"
    assert "Obras" not in (findings[0].evidence_masked or "")


def test_other_path_sources_are_read_and_not_opened(tmp_path):
    path = tmp_path / "fontes.xlsx"
    _book(
        path,
        hyperlink_base="https://empresa.sharepoint.com/sites/obra/arquivo.xlsx",
        template_path=r"C:\Users\fulano\AppData\Roaming\Microsoft\Templates\obra.xltx",
        custom={"Origem": r"C:/Users/fulano/Desktop/nota.xlsx"},
        defined_names=[{"name": "Fora", "formula": r"'C:\Obras\custos.xlsx'!A1"}],
        connections=r"DBQ=C:\Obras\fonte.xlsx;Driver={Microsoft Text Driver};",
    )
    blob = base64.b64encode(b"AAAA" * 20).decode("ascii")
    readable = base64.b64encode(
        "C:/Users/fulano/OneDrive - Empresa/Obra/plan.xlsx".encode("utf-8")
    ).decode("ascii")
    _add_parts(
        path,
        {
            "xl/queryTables/queryTable1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<note url="C:/Obras/consulta.xlsx"/>'
                "</queryTable>"
            ),
            "xl/pivotCache/pivotCacheDefinition1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<cacheSource type="worksheet"><worksheetSource r:id="rId1"/></cacheSource>'
                "</pivotCacheDefinition>"
            ),
            "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
                'Target="\\\\servidor\\dinamica\\fonte.xlsx" TargetMode="External"/>'
                "</Relationships>"
            ),
            "customXml/item1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f"<root><DataMashup>{blob}</DataMashup></root>"
            ),
            "customXml/item2.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f"<root><DataMashup>{readable}</DataMashup></root>"
            ),
        },
    )
    _workbook, findings = _pasta(path)
    by_source = {item.cell: item for item in findings}
    assert by_source["HyperlinkBase"].risk == "alto"
    assert "empresa.sharepoint" in (by_source["HyperlinkBase"].evidence_raw or "")
    assert "sharepoint.com" not in (by_source["HyperlinkBase"].evidence_masked or "")
    assert by_source["Template"].risk == "alto"
    assert "fulano" not in (by_source["Template"].evidence_masked or "")
    assert by_source["custom"].risk == "alto"
    assert by_source["definedName"].risk == "medio"
    assert "custos.xlsx" not in (by_source["definedName"].evidence_masked or "")
    assert by_source["connection"].risk == "medio"
    assert by_source["queryTable"].risk == "medio"
    assert by_source["pivotCache"].risk == "alto"
    assert "servidor" not in (by_source["pivotCache"].evidence_masked or "")
    assert by_source["powerQuery"].risk == "alto"
    assert "Empresa" not in (by_source["powerQuery"].evidence_masked or "")
    assert by_source["powerQuery-presente"].risk == "info"
    assert by_source["powerQuery-presente"].evidence_masked == "Power Query (DataMashup)"
    assert "não foi aberto" in by_source["pivotCache"].message or "não foi executada" in by_source["powerQuery"].message
    hints = [hint for hint in _workbook.network_hints if hint.kind == "pasta"]
    assert {hint.cell for hint in hints} >= set(by_source)


def test_reports_mask_by_default_and_the_map_lists_the_folder(tmp_path, capsys, monkeypatch):
    path = tmp_path / "desktop.xlsx"
    _book(path, abs_path=_ONEDRIVE)
    workbook, findings = _pasta(path)
    text = render_text(findings, show=False, ignored=0, scanned=1, network=workbook.network_hints)
    page = render_html(
        findings, ignored=0, scanned=1, files=[path.name], show=False, network=workbook.network_hints
    )
    raw = render_json(findings, ignored=0, scanned=1, network=workbook.network_hints)
    assert "Pasta onde foi salvo" in text
    assert "Mapa da rede" in text
    assert "fulano" not in text and "Empresa" not in text
    assert "Pasta onde foi salvo" in page and "Mapa da rede" in page
    assert "fulano" not in page and "Empresa" not in page
    assert "fulano" not in raw and "Empresa" not in raw
    assert '"rule": "pasta-salva"' in raw
    revealed = render_html(
        findings, ignored=0, scanned=1, files=[path.name], show=True, network=workbook.network_hints
    )
    assert "fulano" in revealed and "Empresa" in revealed
    monkeypatch.delenv("CI", raising=False)
    monkeypatch.delenv("GITHUB_ACTIONS", raising=False)
    code = main([str(path), "--no-color"])
    out = capsys.readouterr().out
    assert code == 1
    assert "Pasta onde foi salvo" in out
    assert "fulano" not in out
    code = main([str(path), "--show", "--no-color"])
    shown = capsys.readouterr().out
    assert code == 1
    assert "fulano" in shown
    code = main([str(path), "--show", "--format", "json"])
    payload = capsys.readouterr().out
    assert code == 1
    assert "fulano" not in payload and "Empresa" not in payload


def test_window_card_mentions_the_folder_without_the_real_name(tmp_path):
    path = tmp_path / "desktop.xlsx"
    _book(path, abs_path=_LOCAL)
    session = scan_file(path)
    cards = [card for card in window_cards(session, show=False) if card.title == "Pasta onde foi salvo"]
    assert len(cards) == 1
    assert "indício" in cards[0].action
    assert "fulano" not in cards[0].action
    revealed = [card for card in window_cards(session, show=True) if card.title == "Pasta onde foi salvo"]
    assert "fulano" in revealed[0].action


def test_compare_highlights_the_same_user_as_a_clue(tmp_path):
    left = tmp_path / "enviada.xlsx"
    right = tmp_path / "recebida.xlsx"
    _book(left, abs_path=_LOCAL)
    _book(right, abs_path=_ONEDRIVE)
    session = compare_files(left, right)
    change = next(item for item in session.diff.changes if item.type_label == "pasta onde foi salvo")
    assert "mesmo usuário" in change.message
    assert "indício" in change.message
    assert "não prova" in change.message
    assert "fulano" not in change.before_masked and "fulano" not in change.after_masked
    assert "Empresa" not in change.after_masked
    assert "fulano" in change.before_raw and "fulano" in change.after_raw
    assert "Desktop" in change.before_masked
    cards = [card for card in window_cards(session, show=False) if card.title == "Pastas dos dois arquivos"]
    assert cards
    assert "indício" in cards[0].action
    assert "fulano" not in cards[0].action


def test_compare_different_user_company_and_root(tmp_path):
    fulano = tmp_path / "fulano.xlsx"
    beltrano = tmp_path / "beltrano.xlsx"
    ana = tmp_path / "ana.xlsx"
    bia = tmp_path / "bia.xlsx"
    obra_a = tmp_path / "a.xlsx"
    obra_b = tmp_path / "b.xlsx"
    only = tmp_path / "so.xlsx"
    bare = tmp_path / "sem.xlsx"
    _book(fulano, abs_path=_LOCAL)
    _book(beltrano, abs_path="C:/Users/beltrano/Documentos")
    change = next(
        item
        for item in compare_files(fulano, beltrano).diff.changes
        if item.type_label == "pasta onde foi salvo"
    )
    assert "pasta diferente" in change.message
    assert "mesmo usuário" not in change.message
    _book(ana, abs_path="C:/Users/ana/OneDrive - Empresa/Obra/a.xlsx")
    _book(bia, abs_path="C:/Users/bia/OneDrive - Empresa/Outra/b.xlsx")
    change = next(
        item for item in compare_files(ana, bia).diff.changes if item.type_label == "pasta onde foi salvo"
    )
    assert "OneDrive" in change.message or "empresa" in change.message
    assert "indício" in change.message and "não prova" in change.message
    assert "mesmo usuário" not in change.message
    _book(obra_a, abs_path="C:/Obras/lote-a/plan.xlsx")
    _book(obra_b, abs_path="C:/Obras/lote-b/outro.xlsx")
    change = next(
        item
        for item in compare_files(obra_a, obra_b).diff.changes
        if item.type_label == "pasta onde foi salvo"
    )
    assert "pasta raiz" in change.message
    assert "indício" in change.message
    _book(only, abs_path=_LOCAL)
    _book(bare)
    change = next(
        item for item in compare_files(only, bare).diff.changes if item.type_label == "pasta onde foi salvo"
    )
    assert "Só uma" in change.message
    copy = tmp_path / "igual.xlsx"
    shutil.copyfile(fulano, copy)
    same = compare_files(fulano, copy)
    assert same.diff.identical
    assert not any(item.type_label == "pasta onde foi salvo" for item in same.diff.changes)


def test_clean_removes_abs_path_and_other_safe_paths(tmp_path, capsys):
    path = tmp_path / "proposta.xlsx"
    _book(
        path,
        abs_path=_ONEDRIVE,
        hyperlink_base=_LOCAL,
        template_path=r"C:\Modelos\obra.xltx",
        custom={"Origem": "C:/Users/fulano/Desktop/nota.xlsx"},
        defined_names=[
            {"name": "Local", "formula": "Proposta!A1"},
            {"name": "Fora", "formula": r"'C:\Obras\custos.xlsx'!A1"},
        ],
        connections=r"DBQ=C:\Obras\fonte.xlsx;Driver={Microsoft Text Driver};",
    )
    blob = base64.b64encode(b"MASHUP-QUE-FICA-" + b"A" * 80).decode("ascii")
    _add_parts(
        path,
        {
            "xl/queryTables/queryTable1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<queryTable xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
                '<note url="C:/Obras/consulta.xlsx"/>'
                "</queryTable>"
            ),
            "xl/pivotCache/pivotCacheDefinition1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
                'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '<cacheSource type="worksheet"><worksheetSource r:id="rId1"/></cacheSource>'
                "</pivotCacheDefinition>"
            ),
            "xl/pivotCache/_rels/pivotCacheDefinition1.xml.rels": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" '
                'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
                'Target="\\\\servidor\\dinamica\\fonte.xlsx" TargetMode="External"/>'
                "</Relationships>"
            ),
            "customXml/item1.xml": (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                f"<root><DataMashup>{blob}</DataMashup></root>"
            ),
        },
    )
    original = path.read_bytes()
    code = main([str(path), "--limpar", "--no-color"])
    text = capsys.readouterr().out
    dest = tmp_path / "proposta-limpa.xlsx"
    assert code in {0, 1}
    assert path.read_bytes() == original
    assert dest.is_file()
    cleaned = _zip_text(dest)
    assert "absPath" not in cleaned
    assert "AlternateContent" not in cleaned
    assert "fulano" not in cleaned
    assert "Empresa" not in cleaned
    assert "servidor" not in cleaned.casefold()
    assert "consulta.xlsx" not in cleaned
    assert "custos.xlsx" not in cleaned
    assert "fonte.xlsx" not in cleaned
    assert blob in cleaned
    assert "Local" in cleaned
    with zipfile.ZipFile(dest) as archive:
        workbook = archive.read("xl/workbook.xml")
    DefusedET.fromstring(workbook)
    book = open_xlsx(dest)
    assert book["Proposta"]["A1"].value == "Item"
    assert book["Proposta"]["B1"].value == 12
    assert "pasta onde o arquivo foi salvo" in text.casefold() or "caminho" in text.casefold()


def test_libreoffice_opens_a_workbook_without_abs_path(tmp_path):
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        pytest.skip("LibreOffice não está instalado")
    path = tmp_path / "proposta.xlsx"
    _book(path, abs_path=_ONEDRIVE)
    code = main([str(path), "--limpar", "--no-color"])
    assert code in {0, 1}
    dest = tmp_path / "proposta-limpa.xlsx"
    out = tmp_path / "lo"
    out.mkdir()
    env = dict(os.environ)
    env["HOME"] = str(tmp_path / "home")
    (tmp_path / "home").mkdir()
    completed = subprocess.run(
        [soffice, "--headless", "--norestore", "--convert-to", "csv", "--outdir", str(out), str(dest)],
        check=False,
        timeout=120,
        capture_output=True,
        env=env,
    )
    assert completed.returncode == 0, completed.stderr.decode("utf-8", "replace")
    csv_files = list(out.glob("*.csv"))
    assert csv_files
    assert "Item" in csv_files[0].read_text(encoding="utf-8", errors="replace")
